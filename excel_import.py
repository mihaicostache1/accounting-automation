"""Parse an Excel file with the user's existing firms and their declarations.

The spreadsheet layout is whatever the accountant already has, so headers are
matched loosely (diacritics/case ignored): a name column ("nume", "denumire",
"firma"), a CUI column ("cui", "cif", "cod fiscal"), optionally a periodicity
column ("periodicitate", "trimestrial") and one column per declaration type
("100", "D100", "Case de marcat"…). A type cell counts as ticked for anything
except an explicitly empty/negative value, so "x", "da", "1" and even
"trimestrial" all work — and when a 300/394/406 cell *is* "lunar"/"trimestrial"
it also sets the firm's periodicity.

Nothing is written here: `parse_excel` only returns rows with reasons, and the
import dialog shows them for confirmation first — same pattern as the upload
scanner.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import database
from config import DECLARATION_TYPES, PERIODIC_TYPES

_FALSY = {"", "nu", "no", "0", "-", "–", "n", "fals", "false", "none", "nan"}

_NAME_HINTS = ("nume", "denumire", "firma", "societate")
_CUI_HINTS = ("cui", "cif", "cod fiscal", "codfiscal")
_PERIOD_HINTS = ("periodic", "trimestr", "lunar trim")


def _fold(text) -> str:
    """Lowercase, no diacritics, single spaces — for loose header matching."""
    if text is None:
        return ""
    decomposed = unicodedata.normalize("NFKD", str(text))
    ascii_only = "".join(c for c in decomposed if not unicodedata.combining(c))
    return " ".join(re.sub(r"[^0-9a-zA-Z]+", " ", ascii_only).lower().split())


def _match_type_header(folded: str) -> str | None:
    """'100', 'd 100', 'decl 100' -> '100'; anything with 'case' -> Case de marcat."""
    if "case" in folded:
        return "Case de marcat"
    for tip in DECLARATION_TYPES:
        if not tip.isdigit():
            continue
        if re.fullmatch(rf"(d(ecl(aratia|aratie)?)?\s*)?{tip}", folded):
            return tip
    return None


def _periodicity_of(value) -> str | None:
    folded = _fold(value)
    if "trim" in folded:
        return "trimestrial"
    if "lun" in folded:
        return "lunar"
    return None


@dataclass
class ParsedFirm:
    """One spreadsheet row. `error` set means the row cannot be imported;
    `existing` means a firm with this CUI is already in the database."""
    row_number: int
    nume: str = ""
    cui: str = ""
    tipuri: list[str] = field(default_factory=list)
    periodicitate: str = "lunar"
    existing: bool = False
    error: str | None = None


@dataclass
class ParseResult:
    firms: list[ParsedFirm]
    type_columns: list[str]     # declaration types found as columns
    notes: list[str]            # sheet-level remarks for the dialog header


def parse_excel(path: Path | str) -> ParseResult:
    """Read the first sheet and return every data row with a verdict.

    Raises ValueError with a user-facing message when the sheet has no
    recognizable name/CUI header row in its first 10 rows.
    """
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    header_idx, name_col, cui_col = None, None, None
    for i, row in enumerate(rows[:10]):
        folded = [_fold(c) for c in row]
        n = next((j for j, h in enumerate(folded)
                  if any(hint in h for hint in _NAME_HINTS)), None)
        c = next((j for j, h in enumerate(folded)
                  if any(hint in h for hint in _CUI_HINTS)), None)
        if n is not None and c is not None:
            header_idx, name_col, cui_col = i, n, c
            break
    if header_idx is None:
        raise ValueError(
            "Nu am găsit un rând de antet cu coloanele pentru numele firmei "
            "(„Nume”/„Denumire”) și CUI („CUI”/„CIF”) în primele 10 rânduri.")

    header = [_fold(c) for c in rows[header_idx]]
    period_col = next((j for j, h in enumerate(header)
                       if any(hint in h for hint in _PERIOD_HINTS)), None)
    type_cols: dict[int, str] = {}
    for j, h in enumerate(header):
        if j in (name_col, cui_col, period_col) or not h:
            continue
        tip = _match_type_header(h)
        if tip:
            type_cols[j] = tip

    notes = []
    if not type_cols:
        notes.append("Fișierul nu are coloane pe tipuri de declarații — "
                     "firmele importate primesc toate tipurile bifate.")
    if period_col is None and not any(t in PERIODIC_TYPES for t in type_cols.values()):
        notes.append("Fără coloană de periodicitate — se folosește „lunar”.")

    known_cui = {_normalize_cui(f["cui"]) for f in database.list_firme()}
    seen_in_file: set[str] = set()
    firms: list[ParsedFirm] = []
    for i, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if all(_fold(c) == "" for c in row):
            continue
        firm = ParsedFirm(row_number=i)
        firm.nume = str(row[name_col]).strip() if row[name_col] else ""
        raw_cui = str(row[cui_col]).strip() if row[cui_col] else ""
        firm.cui = re.sub(r"\s+", "", raw_cui).upper()
        # Excel loves turning 12345678 into 12345678.0 — undo that.
        firm.cui = re.sub(r"\.0$", "", firm.cui)

        if type_cols:
            for j, tip in type_cols.items():
                value = row[j] if j < len(row) else None
                if _fold(value) in _FALSY:
                    continue
                firm.tipuri.append(tip)
                if tip in PERIODIC_TYPES:
                    if p := _periodicity_of(value):
                        firm.periodicitate = p
            firm.tipuri = [t for t in DECLARATION_TYPES if t in firm.tipuri]
        else:
            firm.tipuri = list(DECLARATION_TYPES)

        if period_col is not None and period_col < len(row):
            if p := _periodicity_of(row[period_col]):
                firm.periodicitate = p

        digits = _normalize_cui(firm.cui)
        if not firm.nume:
            firm.error = "lipsește numele firmei"
        elif not re.fullmatch(r"\d{2,10}", digits):
            firm.error = f"CUI invalid: „{raw_cui or '—'}”"
        elif not firm.tipuri:
            # The firm dialog requires at least one applicable type; importing
            # a firm with none would create an unusable row.
            firm.error = "nicio declarație bifată"
        elif digits in seen_in_file:
            firm.error = "CUI duplicat în fișier"
        elif digits in known_cui:
            firm.existing = True
        seen_in_file.add(digits)
        firms.append(firm)

    return ParseResult(firms=firms,
                       type_columns=[t for t in DECLARATION_TYPES
                                     if t in type_cols.values()],
                       notes=notes)


def _normalize_cui(cui: str) -> str:
    return re.sub(r"\D", "", str(cui))


def import_firms(firms: list[ParsedFirm]) -> int:
    """Insert the given rows (assumed valid and new). Returns how many."""
    count = 0
    for firm in firms:
        firma_id = database.add_firma(firm.nume, firm.cui)
        database.set_firma_tipuri(firma_id, firm.tipuri)
        database.set_firma_periodicitate(firma_id, firm.periodicitate)
        count += 1
    return count
